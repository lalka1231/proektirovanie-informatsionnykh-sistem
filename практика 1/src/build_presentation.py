from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from PIL import Image
P=Path(__file__).resolve().parents[1]; R=P/'diagrams/rendered'
prs=Presentation();prs.core_properties.author="Команда «Байтенки»";prs.core_properties.last_modified_by="Команда «Байтенки»";prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
C={'navy':'20364D','teal':'247E83','mint':'E8F3F4','paper':'F7F9FB','gray':'627387','line':'D9E1E8','amber':'C88C30','pale':'FFF2D7','white':'FFFFFF'}
notes=json.loads((P/'src/notes.json').read_text()) if (P/'src/notes.json').exists() else ['']*15
metadata=[]
def box(sl,x,y,w,h,fill='white',line=None,radius=False):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C.get(fill,fill));sh.line.fill.background() if not line else None
    if line:sh.line.color.rgb=RGBColor.from_string(C.get(line,line))
    if radius:
        try:sh.adjustments[0]=.13
        except:pass
    return sh
def text(sl,txt,x,y,w,h,size=22,color='navy',bold=False,align=None):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(color,color));p.space_after=Pt(6)
        if align is not None:p.alignment=align
    return sh
def arrow(sl,x1,y1,x2,y2,color='teal'):
    ln=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));ln.line.color.rgb=RGBColor.from_string(C[color]);ln.line.width=Pt(2)
    from lxml import etree
    ns='http://schemas.openxmlformats.org/drawingml/2006/main';end=etree.SubElement(ln._element.spPr.find('{'+ns+'}ln'),'{'+ns+'}tailEnd');end.set('type','triangle')
    return ln
def slide(title,kicker='ПРОЕКТИРОВАНИЕ ИНФОРМАЦИОННЫХ СИСТЕМ',source='Учебная модель на основании [1]'):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(C['paper']);n=len(prs.slides)
    box(sl,0,0,.16,7.5,'teal');text(sl,kicker,.5,.25,11,.25,10,'gray',True);text(sl,title,.5,.73,12.25,.72,30,'navy',True)
    box(sl,.5,7.04,12.3,.012,'line');text(sl,source,.5,7.17,11.6,.2,9.5,'gray');text(sl,f'{n:02d} / 15',12,7.14,.83,.25,10,'gray')
    if n<=len(notes):sl.notes_slide.notes_text_frame.text=notes[n-1]
    metadata.append({'slide':n,'title':title,'source':source});return sl

def image(sl,name,x=.55,y=1.65,w=12.2,h=5.12):
    p=R/name;im=Image.open(p);iw,ih=im.size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale
    sl.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
def card(sl,x,y,w,h,head,body,fill='white',big=False):
    box(sl,x,y,w,h,fill,radius=True);text(sl,head,x+.22,y+.2,w-.44,.8,30 if big else 23,'teal',True);text(sl,body,x+.22,y+1.13,w-.44,h-1.25,21)
def badge(sl,txt,x,y,w,kind='pale'):
    box(sl,x,y,w,.42,kind,radius=True);text(sl,txt,x+.12,y+.085,w-.24,.25,12,'gray',True)

def process_slide(title,auto=False):
    sl=slide(title,source='Реконструкция [1]. Укрупнённая схема; подробная Activity Diagram — diagrams/'+('to_be' if auto else 'as_is')+'.puml')
    badge(sl,'ТЕ ЖЕ ВХОДЫ И ГРАНИЦЫ ПРОЦЕССА',.55,1.56,4.6,'mint')
    steps=[('PFT','Задания\nи условия'),('ORION' if auto else 'Планировщик','Проверка\nданных' if auto else 'Проверка\nзаданий'),('ORION' if auto else 'Водитель','Расчет порядка\nостановок' if auto else 'Порядок\nпо опыту'),('Водитель','Доставка\nотправлений'),('Водитель / ИС','Результаты\nвыполнения')]
    x0=.55;w=2.23;gap=.25;y=2.5
    for i,(role,body) in enumerate(steps):
        x=x0+i*(w+gap);text(sl,role,x,y-.38,w,.28,15,'teal' if role=='ORION' else 'gray',True)
        box(sl,x,y,w,1.5,'mint' if role=='ORION' else 'white',radius=True);text(sl,body,x+.12,y+.37,w-.24,1.0,20,'navy',True)
        if i<4:arrow(sl,x+w,y+.74,x+w+gap-.02,y+.74)
    if auto:
        box(sl,.55,4.64,5.65,1.34,'pale',radius=True);text(sl,'Ошибка данных → остановка расчета',.75,4.8,5.2,.38,18,'navy',True);text(sl,'Исправление → новый расчет\nили резервный план',.75,5.24,5.2,.65,18)
        arrow(sl,4.13,4,4.13,4.61,'amber')
        box(sl,6.55,4.64,6.23,1.34,'white',radius=True);text(sl,'Отклонение в пути → разбор с человеком',6.75,4.8,5.8,.38,19,'navy',True);text(sl,'Непрерывный пересчет не приписывается\nпервоначальной версии 2016 года',6.75,5.24,5.8,.65,18)
        arrow(sl,9.12,4,9.12,4.61,'amber')
    else:
        box(sl,.55,4.67,5.65,1.29,'pale',radius=True);text(sl,'Узкое место: ручной выбор порядка',.75,4.85,5.2,.38,20,'navy',True);text(sl,'Не все варианты удается сопоставить',.75,5.32,5.2,.4,19)
        box(sl,6.55,4.67,6.23,1.29,'pale',radius=True);text(sl,'Отклонение → повторное решение',6.75,4.85,5.8,.38,20,'navy',True);text(sl,'Риск лишнего пробега и задержки',6.75,5.32,5.8,.4,19)
        arrow(sl,9.12,4,9.12,4.64,'amber');arrow(sl,6.53,5.32,6.53,4.05,'amber')
    text(sl,'Автоматизировано решение о последовательности. Доставку выполняет человек.' if auto else 'До ORION уже была PFT. Это не переход «с бумаги в компьютер».',.55,6.4,12.1,.38,21,'teal',True)
    return sl

# 1
s=slide('UPS ORION',source='Практическая / самостоятельная работа № 1 • РТУ МИРЭА')
text(s,'Как спроектировать ИС,\nкоторая меняет маршрут доставки',.55,1.73,9,1.55,34,'navy',True)
badge(s,'РЕАЛЬНЫЙ КЕЙС · ВНЕДРЕНИЕ ЗАВЕРШЕНО В 2016 ГОДУ',.55,3.48,9.6,'mint')
for i,(h,b) in enumerate([('БЫЛО','Решение по опыту'),('ИЗМЕНЕНИЕ','Оптимизация + данные'),('СТАЛО','Порядок остановок')]):
    card(s,.55+i*4.12,4.29,3.93,1.85,h,b,'white')
text(s,'Работу выполнила команда «Байтенки»',.55,6.48,11,.34,17,'gray')
# 2
s=slide('Один процесс вместо обзора всей компании',source='Факты о назначении и разработчике: интервью Jack Levis, UPS [1]. Границы — авторские.')
card(s,.55,1.75,3.8,3.25,'UPS','Логистика и доставка\n\nРазработчик ORION —\nсама UPS','white')
card(s,4.6,1.75,3.8,3.25,'Объект анализа','Дневной маршрут\nодного водителя\n\nДоставки + заборы','mint')
card(s,8.65,1.75,4.1,3.25,'Вне границ','Межгород и сортировка\n\nОптимизация всей сети\nи поздние версии','white')
text(s,'Задания дня',.75,5.65,2.8,.55,25,'navy',True);arrow(s,3.4,5.9,4.3,5.9);text(s,'Последовательность',4.6,5.65,3.8,.55,25,'teal',True);arrow(s,8.45,5.9,9.2,5.9);text(s,'Выполнение',9.5,5.65,3,.55,25,'navy',True)
# 3
s=slide('Рост сложности делает опыт недостаточным',source='Менее плотные жилые доставки и запрос на персонализацию — факты [1], Q1. Причинная схема — вывод.')
card(s,.55,1.85,3.85,3.42,'Меньше плотность','Больше расстояние\nмежду остановками','white')
card(s,4.68,1.85,3.85,3.42,'Больше условий','Ожидания клиентов\nи требования сервиса','white')
card(s,8.81,1.85,3.95,3.42,'Сложнее выбор','Снизить пробег,\nсохранив обслуживание','mint')
box(s,.55,5.73,12.2,.7,'navy',radius=True);text(s,'Задача ИС — улучшить решение, а не просто оцифровать работу',.8,5.91,11.65,.39,22,'white',True)
#4
process_slide('AS-IS: порядок уточняет человек',False)
#5
s=slide('Каждой проблеме — проверяемое требование',source='Основания: данные, стабильность маршрута и управление изменениями [1], Q4–Q7. Требования — реконструкция.')
headers=['ПРОБЛЕМА','ТРЕБОВАНИЕ','ЧТО ПРОВЕРЯЕМ']
for x,h in zip([.75,4.8,8.85],headers):text(s,h,x,1.78,3.6,.35,16,'gray',True)
rows=[('Ручное сравнение','Оптимизация с ограничениями','Пробег при том же задании'),('Неточные карты','Контроль качества данных','Пригодность входов'),('Непринятие маршрута','Учет опыта + обучение','Выполнимость и принятие')]
for i,row in enumerate(rows):
    y=2.4+i*1.24;box(s,.55,y,12.2,.97,'white' if i!=1 else 'mint',radius=True)
    for x,txt in zip([.75,4.8,8.85],row):text(s,txt,x,y+.22,3.58,.66,21,'navy',True if x==.75 else False)
text(s,'Нет данных о частоте ошибок — нет вымышленных процентов.',.75,6.43,11,.35,19,'gray')
#6
s=slide('ORION автоматизирует выбор последовательности',source='Собственные PFT и ядро UPS; более 200 000 вариантов за секунды — по интервью [1], Q2–Q3.')
card(s,.55,1.9,3.7,3.04,'Входы','Задания PFT\nКартографические данные\nУсловия обслуживания','white')
card(s,4.65,1.9,3.7,3.04,'Ядро ORION','Математическая\nоптимизация\nс ограничениями','mint')
card(s,8.75,1.9,4,3.04,'Выход','Порядок остановок\nдля водителя\n\nНе робот-доставщик','white')
arrow(s,4.27,3.34,4.6,3.34);arrow(s,8.37,3.34,8.7,3.34)
text(s,'> 200 000',.75,5.42,4,.75,42,'teal',True);text(s,'вариантов маршрута анализируются за секунды',4.9,5.7,7.5,.7,25,'navy',True)
#7
process_slide('TO-BE: система считает, человек выполняет',True)
#8
s=slide('Варианты использования: цели, не код',source='UML Use Case — учебная реконструкция [1]. Подробная спецификация: docs/use_case_specification.md.')
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
def plain_line(sl,x1,y1,x2,y2):
    z=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));z.line.color.rgb=RGBColor.from_string(C['navy']);z.line.width=Pt(1.5);return z
def actor(sl,x,y,label):
    sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-.12),Inches(y-.43),Inches(.24),Inches(.24));sh.fill.background();sh.line.color.rgb=RGBColor.from_string(C['navy'])
    for a,b,c,d in [(x,y-.19,x,y+.2),(x-.25,y-.06,x+.25,y-.06),(x,y+.2,x-.22,y+.44),(x,y+.2,x+.22,y+.44)]:plain_line(sl,a,b,c,d)
    text(sl,label,x-1.04,y+.48,2.08,.59,17,'navy',False,PP_ALIGN.CENTER)
def oval(sl,x,y,w,h,label):
    sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C['mint']);sh.line.color.rgb=RGBColor.from_string(C['teal']);text(sl,label,x+.13,y+.21,w-.26,h-.26,20,'navy',True,PP_ALIGN.CENTER)
box(s,3.05,1.88,7.12,4.82,'white','line');text(s,'ORION — учебная граница системы',3.28,2.08,6.65,.4,22,'teal',True)
actor(s,1.5,2.9,'Планировщик');actor(s,1.5,4.31,'Водитель');actor(s,1.5,5.71,'Руководитель\nдоставки')
actor(s,11.73,2.9,'PFT\nВнешняя ИС');actor(s,11.73,5.14,'Сервис карт\nВнешняя роль')
oval(s,3.7,2.68,2.62,1.03,'Сформировать\nмаршрут');oval(s,3.7,4.07,2.62,1.03,'Получить\nмаршрут');oval(s,3.7,5.47,2.62,1.03,'Оценить\nвыполнение');oval(s,7.06,4.93,2.8,1.24,'Проверить\nисходные данные')
plain_line(s,1.76,2.98,3.69,3.17);plain_line(s,1.76,4.31,3.69,4.57);plain_line(s,1.76,5.71,3.69,5.97)
plain_line(s,11.47,2.92,6.34,3.17);plain_line(s,11.47,5.14,9.88,5.54)
inc=arrow(s,6.32,3.58,7.49,4.92);inc.line.dash_style=MSO_LINE_DASH_STYLE.DASH;text(s,'«include»',6.8,4.12,1.55,.35,16,'teal',True)
#9 custom simplified logical architecture
s=slide('Архитектура: слои известны только логически',source='Реконструированная архитектура на основании доступного описания решения [1]. Физическая топология не установлена.')
badge(s,'РЕКОНСТРУКЦИЯ · МНОГОСЛОЙНАЯ КЛИЕНТ-СЕРВЕРНАЯ МОДЕЛЬ',.55,1.57,10,'pale')
box(s,.65,2.25,8.9,.85,'white',radius=True);text(s,'Клиенты: планировщик · водитель · руководитель',.9,2.48,8.4,.4,23,'navy',True)
box(s,.65,3.57,8.9,1.24,'mint',radius=True);text(s,'Прикладной слой ORION',.9,3.76,8.4,.4,24,'teal',True);text(s,'Проверка входов → собственное ядро UPS → выдача маршрута',.9,4.28,8.4,.35,19)
box(s,.65,5.31,8.9,.95,'white',radius=True);text(s,'Данные: задания · ограничения · версии маршрутов',.9,5.6,8.4,.44,22,'navy',True)
arrow(s,5.1,3.1,5.1,3.53);arrow(s,5.1,4.83,5.1,5.27)
box(s,10.06,2.28,2.69,1.2,'white',radius=True);text(s,'PFT\nПодтверждено [1]',10.24,2.49,2.35,.75,17,'navy',True)
box(s,10.06,4.03,2.69,1.2,'white',radius=True);text(s,'Карты\nПодтверждено [1]',10.24,4.24,2.35,.77,17,'navy',True)
arrow(s,10.03,3.15,9.57,3.88);arrow(s,10.03,4.62,9.57,4.29)
text(s,'Облако? СУБД? Микросервисы? В найденных источниках не установлены.',.75,6.55,11.9,.34,19,'gray')
#10
s=slide('ЖЦ: итерации разработки, внедрение по частям',source='Наиболее подходящая модель ЖЦ, а не установленная методология UPS. Полевые тесты и критерии площадок — факт [1].')
phases=['Анализ\nтребований','Проектирование','Реализация\nи проверки','Внедрение','Эксплуатация']
for i,txt in enumerate(phases):
    x=.55+i*2.48;box(s,x,2.1,2.23,1.46,'mint' if i in [1,2] else 'white',radius=True);text(s,f'0{i+1}',x+.17,2.28,1,.34,17,'teal',True);text(s,txt,x+.1,2.76,2.03,.73,16.5,'navy',True)
    if i<4:arrow(s,x+2.23,2.8,x+2.46,2.8)
arrow(s,6.5,4.05,4.15,4.05);text(s,'Обратная связь полевых тестов',3.15,4.29,5,.4,20,'teal',True)
card(s,.55,5.1,5.88,1.61,'Перед площадкой','Данные и персонал готовы','white');card(s,6.74,5.1,6.01,1.61,'После площадки','Критерии приемки выполнены','white')
#11
s=slide('График проекта: что подтверждено источниками',source='Исторический Гант: [1], [5]. Период и вехи показаны с точностью до года. Учебный график площадки — отдельный файл.')
text(s,'Прототипирование → запуск → распространение решения',.65,1.62,12.0,.43,23,'teal',True)
image(s,'gantt.png',.55,2.28,12.2,4.54)
#12
s=slide('Результат: меньше пробег и стоимость обслуживания',source='Оценки UPS, опубликованные в интервью 2017 года [1]. Эффект ORION отдельно от PFT; не независимый аудит.')
image(s,'results.png',.55,1.65,12.2,4.65)
text(s,'Нет абсолютной базы → нельзя честно показать процент снижения или ROI.',.7,6.42,11.9,.38,21,'gray')
#13
s=slide('Развитие: прогноз времени остановки и его риска',source='Авторское предложение. Не утверждение о текущих возможностях UPS и не обещание измеренного эффекта.')
ph=[('Данные','История остановок\nКонтекст маршрута'),('Машинное обучение','Время +\nнеопределенность'),('Оптимизатор','Резерв времени\nпри высоком риске'),('Человек','Контроль\nисключений')]
for i,(h,b) in enumerate(ph):
    x=.55+i*3.1;card(s,x,2.01,2.88,2.85,h,b,'mint' if i==1 else 'white')
    if i<3:arrow(s,x+2.88,3.52,x+3.06,3.52)
box(s,.55,5.27,12.2,1.17,'pale',radius=True);text(s,'Проверка в пилоте: меньше опозданий без ухудшения пробега и безопасности',.78,5.47,11.7,.4,21,'navy',True);text(s,'При недоверии к модели → базовый расчет. Минимизация персональных данных.',.78,6.0,11.7,.32,18,'gray')
#14
s=slide('Успех ИС = решение + данные + изменение работы',source='Полные URL и реестр фактов — research/sources.md и research/case_analysis.md. Обозначения: факт / вывод / реконструкция.')
for i,(h,b) in enumerate([('Решение','Выбор порядка\nостановок'),('Данные','Точные карты\nи ограничения'),('Люди','Испытания\nи принятие')]):card(s,.55+i*4.14,1.78,3.93,2.4,h,b,'mint' if i==1 else 'white')
text(s,'Опорные источники',.65,4.54,11,.43,24,'navy',True)
a=text(s,'[1] Jack Levis, UPS — интервью ODBMS, 2017',.65,5.14,11.8,.4,21);a.text_frame.paragraphs[0].runs[0].hyperlink.address='https://www.odbms.org/blog/2017/08/big-data-at-ups-interview-with-jack-levis/'
b=text(s,'[4] INFORMS / ORMS Today — «ORION: результат для UPS», 2016',.65,5.72,11.8,.72,20);b.text_frame.paragraphs[0].runs[0].hyperlink.address='https://web.archive.org/web/20250516154643/https://pubsonline.informs.org/do/10.1287/orms.2016.03.10/full/'
text(s,'[4] Архив: 16.05.2025. [5] UPS — пресс-релиз о запуске, 12.04.2016 (GlobeNewswire).',.65,6.52,11.8,.3,15,'gray')
#15
s=slide('Вопросы другим мини-группам',source='Ответы для выступающего и дополнительные вопросы — docs/questions.md.')
qs=['Какое звено AS-IS ограничивало качество решения?','Почему ошибка карты может обесценить хороший алгоритм?','Почему нельзя назвать архитектуру микросервисной по этой схеме?','Что изменить при росте нагрузки в десять раз?','Какой пилот докажет пользу ML, а не только точность прогноза?']
for i,q in enumerate(qs):
    y=1.7+i*.94;box(s,.55,y,.62,.62,'teal',radius=True);text(s,str(i+1),.77,y+.13,.3,.35,21,'white',True);text(s,q,1.4,y+.09,11.25,.65,23,'navy')
text(s,'Главный вывод: проектировать нужно не алгоритм отдельно, а весь процесс принятия решения.',.65,6.62,12,.32,17,'teal',True)
name=P/'Практическая_работа_1_Проектирование_ИС.pptx';prs.save(name)
(P/'src/slides_manifest.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'presentation':str(name),'slides':len(prs.slides),'notes_filled':sum(bool(n) for n in notes),'size':name.stat().st_size},ensure_ascii=False))
