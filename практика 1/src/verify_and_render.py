from pathlib import Path
import pymupdf as fitz
import json,re
from PIL import Image,ImageOps,ImageDraw
from pptx import Presentation
P=Path(__file__).resolve().parents[1];out=P/'preview';out.mkdir(exist_ok=True)
pdf=fitz.open(P/'Практическая_работа_1_Проектирование_ИС.pdf')
for i,page in enumerate(pdf):
    pix=page.get_pixmap(matrix=fitz.Matrix(1.6,1.6));pix.save(out/f'slide_{i+1:02}.png')
for start in range(0,len(pdf),5):
    contact=Image.new('RGB',(1600,5*470),'#dfe5eb');draw=ImageDraw.Draw(contact)
    for j,i in enumerate(range(start,min(start+5,len(pdf)))):
        im=Image.open(out/f'slide_{i+1:02}.png');im.thumbnail((1530,430));contact.paste(im,(40,j*470+30));draw.text((8,j*470+8),str(i+1),fill='#20364d')
    contact.save(out/f'contact_{start+1:02}_{min(start+5,len(pdf)):02}.jpg',quality=92)
prs=Presentation(P/'Практическая_работа_1_Проектирование_ИС.pptx')
errors=[];stats=[]
for i,sl in enumerate(prs.slides,1):
    text=' '.join(s.text for s in sl.shapes if s.has_text_frame)
    for s in sl.shapes:
        if s.left<0 or s.top<0 or s.left+s.width>prs.slide_width+100 or s.top+s.height>prs.slide_height+100:errors.append(f'Slide {i}: out-of-slide shape {s.name}')
    stats.append({'slide':i,'words':len(text.split()),'pictures':sum(s.shape_type==13 for s in sl.shapes),'notes_words':len(sl.notes_slide.notes_text_frame.text.split())})
assert len(prs.slides)==15 and len(pdf)==15
notes=json.loads((P/'src/notes.json').read_text())
assert all(sl.notes_slide.notes_text_frame.text==n for sl,n in zip(prs.slides,notes))
for name in ['as_is','to_be','use_case','architecture','lifecycle','gantt','future_technology','results','gantt_estimate']:
    with Image.open(P/'diagrams/rendered'/f'{name}.png') as im:im.verify()
# Проверка отсечения на уровне PDF: все текстовые spans находятся в области страницы.
for i,page in enumerate(pdf,1):
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines',[]):
            for span in line.get('spans',[]):
                x0,y0,x1,y1=span['bbox']
                if x0<-.5 or y0<-.5 or x1>page.rect.width+.5 or y1>page.rect.height+.5:errors.append(f'PDF page {i}: out-of-page {span["text"]}')
result={'slides':len(prs.slides),'pdf_pages':len(pdf),'speaker_words':sum(len(n.split()) for n in notes),'shape_boundary_errors':errors,'slide_statistics':stats,'verified_diagram_pngs':9,'visual_review':'See preview contact sheets; human/model visual review recorded separately.'}
(P/'docs/verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
assert not errors
